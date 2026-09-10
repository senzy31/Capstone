"""
Renders a ResumeLayout to a DOCX using python-docx. Same layout data as
pdf_generator.py, different renderer - ATS-friendly for the same reasons:
plain paragraphs and built-in list styles, no text boxes, no content in
headers/footers.
"""

from __future__ import annotations

import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor

from app.templates.base import ResumeLayout

ACCENT_COLOR = RGBColor(0x16, 0x4F, 0xAB)
TEXT_COLOR = RGBColor(0x1A, 0x1A, 0x1A)
MUTED_COLOR = RGBColor(0x55, 0x55, 0x55)


def generate_docx(layout: ResumeLayout) -> bytes:
    doc = Document()

    for section in doc.sections:
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = TEXT_COLOR

    _add_name(doc, layout.header.full_name)

    if layout.header.contact_line:
        _add_contact_line(doc, layout.header.contact_line)

    _add_horizontal_rule(doc.add_paragraph(), color="164FAB", size=18)

    for section in layout.sections:
        _add_section_heading(doc, section.heading.upper())

        for entry in section.entries:
            _add_entry(doc, entry)

    buffer = io.BytesIO()
    doc.save(buffer)

    return buffer.getvalue()


def _add_name(doc: Document, name: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = paragraph.add_run(name)
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = TEXT_COLOR

    paragraph.paragraph_format.space_after = Pt(2)


def _add_contact_line(doc: Document, contact_line: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = paragraph.add_run(contact_line)
    run.font.size = Pt(9.5)
    run.font.color.rgb = MUTED_COLOR

    paragraph.paragraph_format.space_after = Pt(10)


def _add_section_heading(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(2)

    run = paragraph.add_run(text)
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = ACCENT_COLOR

    _add_horizontal_rule(doc.add_paragraph(), color="DDDDDD", size=8, space_after=4)


def _add_entry(doc: Document, entry) -> None:
    has_title_row = bool(entry.title or entry.meta)

    if has_title_row:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(0)

        section_width = doc.sections[0].page_width - doc.sections[0].left_margin - doc.sections[0].right_margin
        paragraph.paragraph_format.tab_stops.add_tab_stop(section_width, WD_TAB_ALIGNMENT.RIGHT)

        title_run = paragraph.add_run(entry.title)
        title_run.font.bold = True
        title_run.font.size = Pt(10.5)
        title_run.font.color.rgb = TEXT_COLOR

        if entry.meta:
            paragraph.add_run("\t" + entry.meta).font.color.rgb = MUTED_COLOR
            paragraph.runs[-1].font.size = Pt(9.5)

        if entry.subtitle:
            sub_paragraph = doc.add_paragraph()
            sub_paragraph.paragraph_format.space_after = Pt(2)
            sub_run = sub_paragraph.add_run(entry.subtitle)
            sub_run.font.italic = True
            sub_run.font.size = Pt(10)
            sub_run.font.color.rgb = MUTED_COLOR

    elif entry.subtitle:
        sub_paragraph = doc.add_paragraph()
        sub_run = sub_paragraph.add_run(entry.subtitle)
        sub_run.font.italic = True
        sub_run.font.color.rgb = MUTED_COLOR

    if entry.paragraph:
        p = doc.add_paragraph(entry.paragraph)
        p.paragraph_format.space_after = Pt(2)
        for run in p.runs:
            run.font.size = Pt(10)

    for bullet in entry.bullets:
        p = doc.add_paragraph(bullet, style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        for run in p.runs:
            run.font.size = Pt(10)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def _add_horizontal_rule(paragraph, color: str = "000000", size: int = 6, space_after: int = 8) -> None:
    """python-docx has no native <hr>; a bottom paragraph border is the
    standard way to draw one."""

    paragraph.paragraph_format.space_after = Pt(space_after)

    p_pr = paragraph._p.get_or_add_pPr()
    p_borders = OxmlElement("w:pBdr")

    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)

    p_borders.append(bottom)
    p_pr.append(p_borders)
