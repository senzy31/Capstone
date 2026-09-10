"""
Extracts text from an uploaded PDF or DOCX resume (PyMuPDF / python-docx)
and runs a heuristic structural analysis over it - detected sections,
recognized skills, a keyword list, likely gaps, and basic formatting red
flags. This is pattern-matching over plain text, not a claim of full
resume-parsing accuracy; it's meant to give the user a useful, explainable
report, in the same "don't overclaim" spirit as the rest of this service.
"""

from __future__ import annotations

import io
import re
from typing import List

import fitz  # PyMuPDF
from docx import Document

from app import models
from app.services import ats_service

_SECTION_PATTERNS = {
    "Summary/Objective": [r"\bsummary\b", r"\bobjective\b", r"\bprofile\b"],
    "Experience": [r"\bexperience\b", r"\bemployment\b", r"\bwork history\b"],
    "Education": [r"\beducation\b", r"\bacademic\b"],
    "Skills": [r"\bskills\b", r"\btechnical skills\b", r"\bcompetenc"],
    "Certifications": [r"\bcertification", r"\blicense"],
    "Projects": [r"\bprojects?\b"],
    "Achievements": [r"\bachievements?\b", r"\bawards?\b", r"\bhonors?\b"],
}

_EXPECTED_SECTIONS = ["Summary/Objective", "Experience", "Education", "Skills"]

_EMAIL_PATTERN = re.compile(r"[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}")
_PHONE_PATTERN = re.compile(r"(\+?\d[\d\s().\-]{7,}\d)")


class ParserError(Exception):
    pass


def extract_text_from_pdf(file_bytes: bytes) -> str:
    try:
        with fitz.open(stream=file_bytes, filetype="pdf") as doc:
            return "\n".join(page.get_text() for page in doc)
    except Exception as exc:  # noqa: BLE001 - surfaced as a clean 400 by the route
        raise ParserError(f"Couldn't read this PDF: {exc}") from exc


def extract_text_from_docx(file_bytes: bytes) -> str:
    try:
        document = Document(io.BytesIO(file_bytes))

        paragraphs = [p.text for p in document.paragraphs]

        table_cells = [
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        ]

        return "\n".join(p for p in paragraphs + table_cells if p)
    except Exception as exc:  # noqa: BLE001
        raise ParserError(f"Couldn't read this DOCX file: {exc}") from exc


def analyze_resume_text(text: str) -> models.AnalysisReport:
    lower = text.lower()

    detected_sections = [
        name for name, patterns in _SECTION_PATTERNS.items()
        if any(re.search(p, lower) for p in patterns)
    ]

    missing_sections = [name for name in _EXPECTED_SECTIONS if name not in detected_sections]

    skills_found = sorted({
        term for term in ats_service._SKILL_DICTIONARY
        if ats_service._contains_term(lower, term)
    })

    keywords = ats_service.extract_keywords(text, limit=25)

    experience_found = _excerpt_after_section(text, _SECTION_PATTERNS["Experience"])
    education_found = _excerpt_after_section(text, _SECTION_PATTERNS["Education"])

    formatting_issues = _detect_formatting_issues(text)
    suggestions = _build_suggestions(missing_sections, skills_found, formatting_issues)

    return models.AnalysisReport(
        detected_sections=detected_sections,
        skills_found=skills_found,
        experience_found=experience_found,
        education_found=education_found,
        keywords=keywords,
        missing_sections=missing_sections,
        formatting_issues=formatting_issues,
        suggestions=suggestions,
    )


def _excerpt_after_section(text: str, patterns: List[str], max_lines: int = 6) -> List[str]:
    lines = [line.strip() for line in text.splitlines()]

    for i, line in enumerate(lines):
        if any(re.search(p, line.lower()) for p in patterns) and len(line) < 40:
            excerpt = [l for l in lines[i + 1: i + 1 + max_lines] if l]
            return excerpt

    return []


def _detect_formatting_issues(text: str) -> List[str]:
    issues = []

    word_count = len(text.split())

    if word_count < 60:
        issues.append(
            "Very little text was extracted - if this resume has visible content, it may be "
            "an image-based/scanned PDF, which most ATS systems can't read at all."
        )

    if not _EMAIL_PATTERN.search(text):
        issues.append("No email address was detected - make sure your contact info is in plain text, not an image.")

    if not _PHONE_PATTERN.search(text):
        issues.append("No phone number was detected.")

    tab_heavy_lines = sum(1 for line in text.splitlines() if line.count("\t") >= 3)
    if tab_heavy_lines > 5:
        issues.append(
            "Several lines have heavy tab/column spacing, which often means the resume uses "
            "tables or multi-column layouts - these can scramble the reading order for ATS parsers."
        )

    return issues


def _build_suggestions(missing_sections: List[str], skills_found: List[str], formatting_issues: List[str]) -> List[str]:
    suggestions = []

    for section in missing_sections:
        suggestions.append(f"No clear \"{section}\" section was detected - consider adding one with a distinct heading.")

    if not skills_found:
        suggestions.append(
            "No recognizable technical skills were detected - list specific tools, languages, "
            "or technologies you've actually used so ATS keyword matching can find them."
        )

    if formatting_issues:
        suggestions.append("Review the formatting issues above - they can prevent ATS systems from reading your resume correctly.")

    if not suggestions:
        suggestions.append("No major structural issues detected. Use the ATS Score feature against a specific job description for targeted feedback.")

    return suggestions
